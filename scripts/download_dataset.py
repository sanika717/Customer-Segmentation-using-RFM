import csv, sys
from datetime import datetime
from openpyxl import load_workbook

XLSX = "data/raw/Online Retail.xlsx"
CSV_OUT = "data/raw/online_retail.csv"

print("Converting to CSV (takes 1-3 minutes)...")
wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb.active
n = 0
with open(CSV_OUT, "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f, lineterminator="\n")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            w.writerow([str(c).strip() for c in row])
            continue
        out = []
        for c in row:
            if isinstance(c, datetime):
                c = c.strftime("%Y-%m-%d %H:%M:%S")
            out.append("" if c is None else c)
        w.writerow(out)
        n += 1
        if n % 100000 == 0:
            print(f"  {n} rows...")
wb.close()
print(f"Done: {n} rows written to {CSV_OUT}")
